"""Exports, all generated from the database so they always agree:
calendar.csv (exact 02_ASSETS columns), scheduler_export.csv, per-asset meta/checks/caption, and the workbook."""
from __future__ import annotations

import csv
import json
import os
from datetime import datetime
from pathlib import Path
from zoneinfo import ZoneInfo

from growthengine import __version__
from growthengine.db import DB, loads

CALENDAR_COLUMNS = [
    "asset_id", "platform", "format", "content_pillar", "hook", "caption", "cta", "utm_link", "final_file_url",
    "source_file_url", "source_input", "tools_used", "rights_checked", "ai_label_needed", "scheduled_date",
    "test_published_url", "impressions_or_views", "engagement_rate", "clicks", "status", "notes",
]
SCHEDULER_COLUMNS = ["asset_id", "platform", "publish_at_utc", "publish_at_local", "text", "media_path", "link",
                     "first_comment", "ai_label"]


def _local(cfg, iso: str) -> str:
    if not iso:
        return ""
    tz = ZoneInfo(cfg.schedule.get("timezone", "UTC"))
    dt = datetime.fromisoformat(iso.replace("Z", "+00:00")).astimezone(tz)
    return dt.strftime("%a %d %b %Y %H:%M ") + (dt.tzname() or "")


def source_label(m, source_id: str, title: str) -> str:
    if m is None:
        return source_id
    if m["kind"] == "text":
        return f"{source_id}: section {m['section_idx'] + 1}, paragraph {m['para_idx'] + 1}"
    s, e = m["start_ms"] // 1000, m["end_ms"] // 1000
    return f"{source_id}: {s // 60:02d}:{s % 60:02d}-{e // 60:02d}:{e % 60:02d}"


def tools_label(a, fam: str) -> str:
    writer = "mock writer (deterministic)" if a["llm"] == "mock" else f"Claude API ({a['llm']})"
    renderer = {"video": "ffmpeg", "carousel": "Pillow", "graphic": "Pillow"}.get(fam, "text")
    return f"growthengine {__version__}; {writer}; {renderer}"


def calendar_rows(settings, db: DB) -> list[dict]:
    cfg = settings.config
    base = os.environ.get("GE_CONTENT_BASE_URL", "").rstrip("/")
    rows = []
    assets = db.q("SELECT a.*, i.ai_media, i.title AS input_title FROM assets a "
                  "LEFT JOIN inputs i ON i.input_key = a.input_key "
                  "WHERE a.status = 'ready' ORDER BY a.scheduled_at, a.asset_id")
    for a in assets:
        m = db.one("SELECT * FROM moments WHERE key=?", (a["moment_key"],))
        fam = cfg.formats[a["format"]]["family"]
        files = loads(a["files_json"], {})
        copy = loads(a["copy_json"], {})
        pillar = cfg.brand["pillars"].get(a["pillar"], {}).get("name", "")
        prefix = f"{base}/" if base else ""
        target = cfg.mix.get("targets", {}).get(a["format"], "")
        notes = [f"target: {target}", f"slot: {_local(cfg, a['scheduled_at'])}",
                 f"link: {cfg.platforms[a['platform']].get('link_mode')}", f"v{a['version']}"]
        if copy.get("title"):
            notes.insert(0, f"YouTube title: {copy['title']}")
        if a["ai_media"]:
            notes.append("apply the platform's AI label when posting")
        rows.append({
            "asset_id": a["asset_id"], "platform": a["platform"], "format": a["format"],
            "content_pillar": f"{a['pillar']} {pillar}".strip(), "hook": copy.get("hook", ""),
            "caption": a["caption"], "cta": a["cta"], "utm_link": a["utm_link"],
            "final_file_url": f"{prefix}assets/{a['asset_id']}/{files.get('final', '')}",
            "source_file_url": f"{prefix}assets/{a['asset_id']}/{files.get('source', '')}",
            "source_input": source_label(m, a["source_id"], a["input_title"] or ""),
            "tools_used": tools_label(a, fam), "rights_checked": "Yes",
            "ai_label_needed": "Yes" if a["ai_media"] else "No", "scheduled_date": a["scheduled_at"],
            "test_published_url": a["test_url"] or "", "impressions_or_views": a["impressions"] or "",
            "engagement_rate": a["engagement_rate"] or "", "clicks": a["clicks"] or "",
            "status": "Ready", "notes": "; ".join(notes),
        })
    return rows


def write_csv(path: Path, columns: list[str], rows: list[dict]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        w.writerows(rows)


def scheduler_rows(settings, db: DB, cal: list[dict]) -> list[dict]:
    out = []
    for r in cal:
        a = db.one("SELECT files_json FROM assets WHERE asset_id=?", (r["asset_id"],))
        files = loads(a["files_json"], {})
        media = files.get("final", "")
        mode = settings.config.platforms[r["platform"]].get("link_mode")
        out.append({
            "asset_id": r["asset_id"], "platform": r["platform"], "publish_at_utc": r["scheduled_date"],
            "publish_at_local": _local(settings.config, r["scheduled_date"]), "text": r["caption"],
            "media_path": "" if media.endswith(".txt") else f"assets/{r['asset_id']}/{media}",
            "link": r["utm_link"], "first_comment": r["utm_link"] if mode == "first_comment" else "",
            "ai_label": r["ai_label_needed"],
        })
    return out


def write_asset_sidecars(settings, db: DB, asset_id: str) -> None:
    a = db.one("SELECT * FROM assets WHERE asset_id=?", (asset_id,))
    out = settings.content_dir / "assets" / asset_id
    out.mkdir(parents=True, exist_ok=True)
    m = db.one("SELECT * FROM moments WHERE key=?", (a["moment_key"],))
    i = db.one("SELECT * FROM inputs WHERE input_key=?", (a["input_key"],))
    (out / "caption.txt").write_text(a["caption"] + "\n", encoding="utf-8")
    (out / "checks.json").write_text(a["checks_json"] or "[]", encoding="utf-8")
    meta = {k: a[k] for k in ("asset_id", "platform", "format", "pillar", "status", "cta", "utm_link",
                              "scheduled_at", "version", "created_run", "updated_run", "llm")}
    meta.update({
        "copy": loads(a["copy_json"], {}), "files": loads(a["files_json"], {}), "fonts": loads(a["fonts_json"], []),
        "source": {"source_id": a["source_id"], "input": i["path"] if i else "", "title": i["title"] if i else "",
                   "rights": i["rights_note"] if i else "",
                   "location": source_label(m, a["source_id"], ""),
                   "quote": m["quote"] if m else "", "point": m["point"] if m else ""},
    })
    (out / "meta.json").write_text(json.dumps(meta, indent=2, ensure_ascii=False), encoding="utf-8")


def export_all(settings, db: DB) -> dict:
    cal = calendar_rows(settings, db)
    write_csv(settings.content_dir / "calendar.csv", CALENDAR_COLUMNS, cal)
    write_csv(settings.content_dir / "scheduler_export.csv", SCHEDULER_COLUMNS, scheduler_rows(settings, db, cal))
    for r in cal:
        write_asset_sidecars(settings, db, r["asset_id"])
    return {"calendar": settings.content_dir / "calendar.csv", "rows": len(cal)}


def fill_workbook(settings, db: DB, src: Path, dst: Path) -> int:
    """Write calendar rows into 02_ASSETS columns A-U from row 5 of a COPY of the official workbook.
    Formulas (V/W and every other tab) are never touched."""
    from openpyxl import load_workbook

    if src.resolve() == dst.resolve():
        raise SystemExit("refusing to overwrite the source workbook; pass a different --out")
    wb = load_workbook(src)
    ws = wb["02_ASSETS"]
    header = [ws.cell(row=4, column=c).value for c in range(1, len(CALENDAR_COLUMNS) + 1)]
    if header != CALENDAR_COLUMNS:
        raise SystemExit(f"02_ASSETS header does not match calendar.csv columns: {header}")
    rows = calendar_rows(settings, db)
    last = 154
    if len(rows) > last - 4:
        raise SystemExit(f"{len(rows)} assets do not fit in 02_ASSETS rows 5-{last}")
    for r in range(5, last + 1):
        for c in range(1, len(CALENDAR_COLUMNS) + 1):
            ws.cell(row=r, column=c).value = None
    for i, row in enumerate(rows):
        for c, col in enumerate(CALENDAR_COLUMNS, 1):
            v = row[col]
            if col in ("impressions_or_views", "clicks") and str(v).isdigit():
                v = int(v)
            ws.cell(row=5 + i, column=c).value = v if v != "" else None
    dst.parent.mkdir(parents=True, exist_ok=True)
    wb.save(dst)
    return len(rows)
