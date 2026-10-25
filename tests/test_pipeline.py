"""End to end with the mock writer: produce, repeat, refresh, fail + resume, export, audit."""
import csv
import json
import shutil
from pathlib import Path

import pytest

from conftest import ROOT, run
from growthengine.db import DB
from growthengine.export import CALENDAR_COLUMNS, fill_workbook
from growthengine.pipeline import RunFailed

DROPDOWNS = json.loads((Path(__file__).parent / "fixtures" / "assets_dropdowns.json").read_text(encoding="utf-8"))


def calendar(settings):
    with (settings.content_dir / "calendar.csv").open(encoding="utf-8-sig", newline="") as f:
        return list(csv.DictReader(f))


def test_one_input_makes_a_week_of_ready_assets(settings):
    r = run(settings, "samples/I2_blog.md")
    assert r["status"] == "done"
    assert r["produced"] >= 10
    rows = calendar(settings)
    assert len(rows) == r["library_ready"] == r["produced"]
    assert list(rows[0].keys()) == CALENDAR_COLUMNS
    allowed = {col: set(v.strip('"').split(",")) for col, v in
               zip(["platform", "format", "rights_checked", "ai_label_needed", "status"],
                   [DROPDOWNS["B5:B154"], DROPDOWNS["C5:C154"], DROPDOWNS["M5:M154"], DROPDOWNS["N5:N154"],
                    DROPDOWNS["T5:T154"]])}
    for row in rows:
        for col, ok in allowed.items():
            assert row[col] in ok, (col, row[col])
        assert row["utm_link"].endswith("utm_content=" + row["asset_id"].lower())
        assert (settings.content_dir / row["final_file_url"].rstrip("/")).exists()
        assert (settings.content_dir / row["source_file_url"].rstrip("/")).exists()
    assert len({r["platform"] for r in rows}) >= 3
    assert len({r["scheduled_date"] + r["platform"] for r in rows}) == len(rows)   # no double-booked slot


def test_repeat_run_makes_nothing_new_and_costs_nothing(settings):
    first = run(settings, "samples/I2_blog.md")
    again = run(settings, "samples/I2_blog.md")
    assert again["produced"] == 0 and again["calls"] == 0
    assert again["unchanged"] == first["produced"]
    assert len(calendar(settings)) == first["produced"]


def test_refresh_adds_only_the_new_section_and_keeps_ids(settings):
    run(settings, "samples/I2_blog.md")
    before = {r["asset_id"]: r["scheduled_date"] for r in calendar(settings)}
    r = run(settings, "samples/I2_blog_v2.md")
    after = {r_["asset_id"]: r_["scheduled_date"] for r_ in calendar(settings)}
    assert r["produced"] >= 1 and r["duplicates"] == 0
    assert all(after[k] == v for k, v in before.items())        # same ids, same slots
    new = set(after) - set(before)
    v2 = (settings.root / "samples" / "I2_blog_v2.md").read_text(encoding="utf-8")
    new_section = v2.split("## Batch the busywork into one sitting")[1].split("## ")[0]
    db = DB(settings.db_path)
    for aid in new:
        m = db.one("SELECT m.* FROM assets a JOIN moments m ON m.key=a.moment_key WHERE a.asset_id=?", (aid,))
        assert m["quote"] in new_section


def test_failure_then_resume_matches_a_clean_run(settings, tmp_path_factory):
    settings.fault = "llm_timeout@write:3"
    with pytest.raises(RunFailed) as e:
        run(settings, "samples/I2_blog.md")
    assert e.value.stage == "write"
    resumed = run(settings, None, resume=e.value.run_id)
    assert resumed["status"] == "done"
    ids_resumed = sorted(r["asset_id"] for r in calendar(settings))

    # same input in a fresh project, no fault
    from growthengine.settings import load_settings
    clean_root = tmp_path_factory.mktemp("clean")
    shutil.copytree(ROOT / "config", clean_root / "config")
    shutil.copytree(settings.root / "samples", clean_root / "samples")
    clean = load_settings(clean_root)
    run(clean, str(clean_root / "samples" / "I2_blog.md"))
    assert ids_resumed == sorted(r["asset_id"] for r in calendar(clean))
    log = (settings.log_dir / f"run-{e.value.run_id}.jsonl").read_text(encoding="utf-8")
    assert '"level": "error"' in log and "injected fault" in log


def test_same_text_under_another_name_never_duplicates_an_asset(settings, project):
    run(settings, "samples/I2_blog.md")
    shutil.copy(project / "samples" / "I2_blog.md", project / "samples" / "copy_of_blog.md")
    run(settings, "samples/copy_of_blog.md")
    # identical text -> identical moments. An idea may get its second platform, but never a second copy of an
    # asset, never two assets on one platform, never more than max_assets_per_moment.
    db = DB(settings.db_path)
    rows = db.q("SELECT moment_key, platform, format FROM assets WHERE status='ready'")
    per_moment = {}
    for r in rows:
        per_moment.setdefault(r["moment_key"], []).append((r["platform"], r["format"]))
    for uses in per_moment.values():
        assert len(uses) <= settings.config.mix["max_assets_per_moment"]
        assert len({p for p, _ in uses}) == len(uses) and len({f for _, f in uses}) == len(uses)
    assert len(calendar(settings)) == len({x["asset_id"] for x in calendar(settings)})


def test_near_duplicate_detection_is_per_platform():
    from growthengine.dedupe import near_duplicate
    lib = [{"asset_id": "LI-TX-1", "platform": "LinkedIn", "hook": "Cut clips from the middle of your episode",
            "body": "The best moments sit in the middle, where guests relax."}]
    assert near_duplicate("LinkedIn", "Cut clips from the middle of the episode", "", lib)[0] == "LI-TX-1"
    assert near_duplicate("X", "Cut clips from the middle of the episode", "", lib) is None
    assert near_duplicate("LinkedIn", "Plan the week before you post", "Something else entirely here.", lib) is None


def test_audit_passes_and_catches_tampering(settings):
    from growthengine.audit import audit
    run(settings, "samples/I2_blog.md")
    db = DB(settings.db_path)
    res = audit(settings, db, 20, seed=1)
    assert res["rate"] == 1.0
    victim = res["sample"][0]
    (settings.content_dir / "assets" / victim / "caption.txt").write_text("This is a game-changer\n", encoding="utf-8")
    res2 = audit(settings, db, 20, seed=1)
    assert victim in res2["failures"]


def test_human_reject_removes_asset_from_main_output(settings):
    from growthengine.cli import main
    run(settings, "samples/I2_blog.md")
    victim = calendar(settings)[0]["asset_id"]
    assert main(["reject", victim, "--reason", "hook is weak"]) == 0
    assert victim not in {r["asset_id"] for r in calendar(settings)}
    assert not (settings.content_dir / "assets" / victim).exists()
    assert (settings.out_dir / "rejected" / victim).exists()


@pytest.mark.skipif(not (ROOT / "brief" / "workbook_original.xlsx").is_file(),
                    reason="confidential workbook not present (kept out of the repo)")
def test_workbook_fill_keeps_formulas(settings, tmp_path):
    from openpyxl import load_workbook
    run(settings, "samples/I2_blog.md")
    out = tmp_path / "filled.xlsx"
    n = fill_workbook(settings, DB(settings.db_path), ROOT / "brief" / "workbook_original.xlsx", out)
    ws = load_workbook(out)["02_ASSETS"]
    assert ws["A5"].value and ws.cell(row=4 + n, column=1).value and ws.cell(row=5 + n, column=1).value is None
    assert str(ws["V5"].value).startswith("=IF(") and str(ws["W5"].value).startswith("=IF(")
    assert str(load_workbook(out)["04_SELF_CHECK"]["D6"].value).startswith("=COUNTIF(")
